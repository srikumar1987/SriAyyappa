"""Generates the store data workbook. Upload to Google Drive, then open with Google Sheets.

The column names here are the contract with src/data/repo.ts. Do not rename them.
Run:  python scripts/make_sheet_template.py
Output: data/Ayyappa_Pooja_Store_Data.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "data" / "Ayyappa_Pooja_Store_Data.xlsx"
MAX_ROW = 501  # formulas and validation cover rows 2..501

HEADER_FILL = PatternFill("solid", fgColor="FF9933")
HEADER_FONT = Font(bold=True, color="FFFFFF")

CATEGORIES = ["Pooja Items", "Idols", "Incense", "Kerala Snacks", "Kerala Sweets", "Kerala Products"]

# Same 16 products as src/data/products.ts
PRODUCTS = [
    ("P001", "പിത്തള വിളക്ക്", "Brass Lamp", "பித்தளை விளக்கு", "ಹಿತ್ತಾಳೆ ದೀಪ", "ఇత్తడి దీపం", "पीतल का दीया", "Pooja Items", 200),
    ("P002", "അഗർബത്തി", "Agarbatti", "அகர்பத்தி", "ಅಗರ್ಬತ್ತಿ", "అగరబత్తి", "अगरबत्ती", "Incense", 15),
    ("P003", "പൂജാ താലം", "Pooja Thali", "பூஜை தட்டு", "ಪೂಜಾ ತಟ್ಟೆ", "పూజా తాళం", "पूजा थाली", "Pooja Items", 250),
    ("P004", "കുങ്കുമം", "Kumkum", "குங்குமம்", "ಕುಂಕುಮ", "కుంకుమ", "कुंकुम", "Pooja Items", 30),
    ("P005", "കർപ്പൂരം", "Camphor", "கற்பூரம்", "ಕರ್ಪೂರ", "కర్పూరం", "कपूर", "Pooja Items", 20),
    ("P006", "പിത്തള വിഗ്രഹം", "Brass Idol", "பித்தளை சிலை", "ಹಿತ್ತಾಳೆ ವಿಗ್ರಹ", "ఇత్తడి విగ్రహం", "पीतल की मूर्ति", "Idols", 500),
    ("K001", "ഉപ്പേരി", "Banana Chips", "வாழைக்காய் சிப்ஸ்", "ಬಾಳೆಕಾಯಿ ಚಿಪ್ಸ್", "అరటికాయ చిప్స్", "केले के चिप्स", "Kerala Snacks", 80),
    ("K002", "മിക്സ്ചർ", "Mixture", "மிக்சர்", "ಮಿಕ್ಸ್ಚರ್", "మిక్చర్", "मिक्सचर", "Kerala Snacks", 120),
    ("K003", "മുറുക്ക്", "Murukku", "முறுக்கு", "ಮುರುಕ್ಕು", "మురుక్కు", "मुरुक्कु", "Kerala Snacks", 100),
    ("K004", "അച്ചപ്പം", "Achappam", "அச்சப்பம்", "ಅಚ್ಚಪ್ಪಂ", "అచ్చప్పం", "अच्चप्पम", "Kerala Snacks", 90),
    ("K005", "ചക്ക ചിപ്സ്", "Jackfruit Chips", "பலா சிப்ஸ்", "ಹಲಸಿನ ಹಣ್ಣಿನ ಚಿಪ್ಸ್", "పనస చిప్స్", "कटहल के चिप्स", "Kerala Snacks", 150),
    ("K006", "കപ്പ ചിപ്സ്", "Tapioca Chips", "மரவள்ளிக்கிழங்கு சிப்ஸ்", "ಮರಗೆಣಸು ಚಿಪ್ಸ್", "కర్రపెండలం చిప్స్", "कसावा चिप्स", "Kerala Snacks", 70),
    ("K007", "പായസം മിക്സ്", "Payasam Mix", "பாயசம் மிக்ஸ்", "ಪಾಯಸ ಮಿಕ್ಸ್", "పాయసం మిక్స్", "पायसम मिक्स", "Kerala Sweets", 180),
    ("K008", "വെളിച്ചെണ്ണ", "Coconut Oil", "தேங்காய் எண்ணெய்", "ತೆಂಗಿನೆಣ್ಣೆ", "కొబ్బరి నూనె", "नारियल तेल", "Kerala Products", 250),
    ("K009", "കേരള മസാല പായ്ക്ക്", "Kerala Spices Pack", "கேரள மசாலா தொகுப்பு", "ಕೇರಳ ಮಸಾಲಾ ಪ್ಯಾಕ್", "కేరళ మసాలా ప్యాక్", "केरल मसाला पैक", "Kerala Products", 350),
    ("K010", "ശർക്കര", "Jaggery", "வெல்லம்", "ಬೆಲ್ಲ", "బెల్లం", "गुड़", "Kerala Products", 60),
]

# Sheet name -> (headers, column widths). Formula columns are added in add_formulas().
SCHEMAS = {
    "Products": (
        ["Product_ID", "Name_Malayalam", "Name_English", "Name_Tamil", "Name_Kannada",
         "Name_Telugu", "Name_Hindi", "Category", "Default_Selling_Price", "Image_Folder"],
        [12, 26, 24, 26, 26, 26, 24, 18, 20, 20],
    ),
    "Lots": (
        ["Lot_ID", "Product_ID", "Received_On", "Supplier", "Buy_Price_Unit", "Qty_Received",
         "Selling_Price_Unit", "Expense_Note",
         "Lot_Expenses_Total", "Qty_Sold", "Qty_Remaining", "Landed_Cost_Unit", "Status"],
        [12, 12, 14, 22, 16, 16, 18, 26, 20, 12, 16, 18, 12],
    ),
    "Sales": (
        ["Sale_ID", "Slice_No", "Sold_At", "Product_ID", "Lot_ID", "Qty", "Unit_Price", "Unit_Cost",
         "Payment", "Customer_Name", "Customer_Phone", "Total"],
        [12, 10, 22, 12, 12, 10, 14, 14, 14, 22, 16, 14],
    ),
    "Expenses": (
        ["Expense_ID", "Lot_ID", "Type", "Amount", "Date", "Note"],
        [12, 12, 20, 14, 14, 36],
    ),
    "Overhead": (
        ["Overhead_ID", "Month", "Type", "Amount", "Note"],
        [14, 12, 16, 14, 36],
    ),
    "Quotes": (
        ["Quote_ID", "Created_At", "Customer_Name", "Customer_Phone", "Language", "Status", "Total", "Items_JSON"],
        [12, 22, 22, 16, 10, 12, 12, 60],
    ),
    "Stock": (
        ["Product_ID", "Product Name (English)", "Total Received", "Total Sold", "Remaining",
         "Last Sale Date", "Days Since Last Sale", "Stock Status"],
        [12, 26, 16, 14, 14, 16, 18, 18],
    ),
}

DROPDOWNS = {
    ("Products", "H"): CATEGORIES,
    ("Expenses", "C"): ["Transport", "Packaging", "Loading/Unloading", "Tax", "Other"],
    ("Overhead", "C"): ["rent", "electricity", "staff", "other"],
    ("Sales", "I"): ["cash", "upi", "whatsapp", "bank"],
    ("Quotes", "F"): ["draft", "sent", "accepted", "expired"],
}


def style_header(ws, headers, widths):
    for col, (title, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"


def add_formulas(wb):
    lots = wb["Lots"]
    for r in range(2, MAX_ROW + 1):
        lots[f"I{r}"] = f'=IF(A{r}="","",SUMIFS(Expenses!D:D,Expenses!B:B,A{r}))'
        lots[f"J{r}"] = f'=IF(A{r}="","",SUMIFS(Sales!F:F,Sales!E:E,A{r}))'
        lots[f"K{r}"] = f'=IF(A{r}="","",F{r}-J{r})'
        lots[f"L{r}"] = f'=IF(A{r}="","",IF(F{r}=0,"",(E{r}*F{r}+I{r})/F{r}))'
        lots[f"M{r}"] = f'=IF(A{r}="","",IF(K{r}<=0,"Sold Out","Active"))'

    sales = wb["Sales"]
    for r in range(2, MAX_ROW + 1):
        sales[f"L{r}"] = f'=IF(A{r}="","",F{r}*G{r})'

    stock = wb["Stock"]
    for r in range(2, MAX_ROW + 1):
        stock[f"A{r}"] = f'=IF(Products!A{r}="","",Products!A{r})'
        stock[f"B{r}"] = f'=IF(A{r}="","",Products!C{r})'
        stock[f"C{r}"] = f'=IF(A{r}="","",SUMIFS(Lots!F:F,Lots!B:B,A{r}))'
        stock[f"D{r}"] = f'=IF(A{r}="","",SUMIFS(Sales!F:F,Sales!D:D,A{r}))'
        stock[f"E{r}"] = f'=IF(A{r}="","",C{r}-D{r})'
        stock[f"F{r}"] = f'=IF(A{r}="","",IF(D{r}=0,"",MAXIFS(Sales!C:C,Sales!D:D,A{r})))'
        stock[f"G{r}"] = f'=IF(A{r}="","",IF(F{r}="","",TODAY()-F{r}))'
        stock[f"H{r}"] = (
            f'=IF(A{r}="","",IF(E{r}<=0,"Out of stock",IF(E{r}<=5,"Low stock",'
            f'IF(AND(G{r}<>"",G{r}>=30),"Slow moving","OK"))))'
        )
        stock[f"F{r}"].number_format = "DD-MMM-YYYY"


def add_validation(wb):
    for (sheet, col), options in DROPDOWNS.items():
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        wb[sheet].add_data_validation(dv)
        dv.add(f"{col}2:{col}{MAX_ROW}")


def add_readme(wb):
    ws = wb.create_sheet("Read Me", 0)
    lines = [
        "Ayyappa Pooja Store - Data Workbook",
        "",
        "Products: product list (names in 6 languages). Add a row for each new product.",
        "Lots: one row per stock arrival. Lot_Expenses_Total and stock columns fill themselves in.",
        "Expenses: extra costs for a lot (transport, packing...). Use the same Lot_ID as the Lots row.",
        "Sales: one row per FIFO slice of a sale. The website writes these rows; do not type sales here.",
        "Overhead: monthly running costs. Month is YYYY-MM.",
        "Quotes: quotations made on the website. Change Status as they progress.",
        "Stock: automatic. Do not type here.",
        "",
        "Keep sheet names and column headings exactly as they are.",
        "",
        "ഉൽപ്പന്നങ്ങൾ: ഓരോ പുതിയ ഉൽപ്പന്നത്തിനും ഒരു വരി ചേർക്കുക.",
        "ലോട്ടുകൾ: പുതിയ സ്റ്റോക്ക് വരുമ്പോൾ ഒരു വരി ചേർക്കുക.",
        "ചെലവുകൾ: ലോട്ടിനുള്ള അധിക ചെലവുകൾ. അതേ Lot_ID ഉപയോഗിക്കുക.",
        "സ്റ്റോക്ക്: ബാക്കിയുള്ളവയുടെ പട്ടിക. ഇതിൽ ഒന്നും ടൈപ്പ് ചെയ്യേണ്ടതില്ല.",
    ]
    for i, text in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=text)
    ws["A1"].font = Font(bold=True, size=14, color="C41E3A")
    ws.column_dimensions["A"].width = 120


def main():
    wb = Workbook()
    wb.remove(wb.active)

    for name, (headers, widths) in SCHEMAS.items():
        style_header(wb.create_sheet(name), headers, widths)

    prod = wb["Products"]
    for r, row in enumerate(PRODUCTS, start=2):
        for c, value in enumerate(row, start=1):
            prod.cell(row=r, column=c, value=value)
        prod.cell(row=r, column=10, value=f"{row[0]}_{row[2].replace(' ', '_')}")

    add_formulas(wb)
    add_validation(wb)
    add_readme(wb)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
